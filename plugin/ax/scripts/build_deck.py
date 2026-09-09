#!/usr/bin/env python3
"""build_deck.py — proposal/pages/*.md (v3 propdraft) → 발표자료 PPTX (결정론 빌더, A3 부수효과).

toc.json 순서로 페이지를 정렬해 슬라이드 생성. 각 슬라이드 = 제목(장·절) + governance(핵심 메시지) + 본문 개요.
usage: build_deck.py [proposalDir=proposal] [out=proposal/build/발표자료.pptx]
의존: python-pptx, PyYAML.
"""
import sys, re, json, pathlib
try:
    import yaml
    from pptx import Presentation
    from pptx.util import Pt, Emu
except Exception as e:
    print(f"[build_deck] 의존성 필요(python-pptx·PyYAML): {e}"); sys.exit(1)

PROP = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "proposal")
OUT = pathlib.Path(sys.argv[2] if len(sys.argv) > 2 else PROP / "build" / "발표자료.pptx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FM = re.compile(r'^---\n(.*?)\n---\n(.*)$', re.DOTALL)

def parse_page(md):
    m = FM.match(md)
    if not m:
        return None
    fmy = yaml.safe_load(m.group(1)) or {}
    pd = fmy.get('propdraft', {})
    body = m.group(2).strip()
    return pd, body

def page_order():
    toc = json.loads((PROP / "toc.json").read_text(encoding="utf-8")) if (PROP / "toc.json").exists() else {}
    ids = [pg["page_id"] for ch in toc.get("chapters", []) for se in ch.get("sections", []) for pg in se.get("pages", [])]
    return ids

def add_slide(prs, pd, body):
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    title = f"{pd.get('chapter_number','')}. {pd.get('chapter_title','')} — {pd.get('section_label','')}"
    slide.shapes.title.text = title.strip(" —.")
    tf = slide.placeholders[1].text_frame
    tf.clear()
    sub = (pd.get('subtitle') or {}).get('text')
    if sub:
        p = tf.paragraphs[0]; p.text = f"[{sub}]"; p.font.bold = True; p.font.size = Pt(16)
    lead = re.compile(r'^[○\-\s]+')
    for g in (pd.get('governance') or [])[:5]:
        gtext = lead.sub('', str(g))
        p = tf.add_paragraph(); p.text = "○ " + gtext; p.level = 0; p.font.size = Pt(14)
    # 본문 첫 불릿 2~3개 요약
    for line in [l for l in body.splitlines() if l.strip().startswith(('-', '·'))][:3]:
        p = tf.add_paragraph(); p.text = "  - " + line.strip(' -·')[:80]; p.level = 1; p.font.size = Pt(11)

def main():
    prs = Presentation()
    # 표지
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = "제안 발표자료 (자동 생성)"
    try:
        s0.placeholders[1].text = "ax 제안축 파이프라인 · propdraft → PPTX"
    except KeyError:
        pass
    order = page_order()
    pages = {p.stem: p for p in (PROP / "pages").glob("*.md")} if (PROP / "pages").exists() else {}
    ordered = [pid for pid in order if pid in pages] + [k for k in pages if k not in order]
    n = 0
    for pid in ordered:
        parsed = parse_page(pages[pid].read_text(encoding="utf-8"))
        if not parsed:
            print(f"[build_deck] ⚠ {pid} frontmatter 무효 — 스킵"); continue
        add_slide(prs, *parsed); n += 1
    prs.save(OUT)
    print(f"[build_deck] 슬라이드 {n+1}장(표지+{n}) → {OUT}")

if __name__ == "__main__":
    main()
